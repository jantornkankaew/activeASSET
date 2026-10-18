#!/usr/bin/env python3
"""
activeASSET — Import Demo Assets from Excel (Python version)
อ่าน /home/z/my-project/upload/Fixed Asset_A5.xlsx แล้ว import เข้า SQLite DB ตรงๆ
"""
import openpyxl
import sqlite3
import json
import os
import sys
from datetime import datetime

EXCEL_PATH = "/home/z/my-project/upload/Fixed Asset_A5.xlsx"
DB_PATH = "/home/z/my-project/db/activeasset.db"
TENANT_CODE = "AIS"

# === Find DB path ===
if not os.path.exists(DB_PATH):
    # หาจาก .env
    env_path = "/home/z/my-project/.env"
    if os.path.exists(env_path):
        with open(env_path) as f:
            for line in f:
                if line.startswith("DATABASE_URL="):
                    url = line.strip().split("=", 1)[1].strip().strip('"').strip("'")
                    if "file:" in url:
                        DB_PATH = url.split("file:")[-1]
                        if not os.path.isabs(DB_PATH):
                            DB_PATH = os.path.join("/home/z/my-project", DB_PATH)
                    break

print(f"📂 Excel: {EXCEL_PATH}")
print(f"🗄️  DB:    {DB_PATH}")

if not os.path.exists(DB_PATH):
    print(f"❌ DB not found at {DB_PATH}")
    sys.exit(1)

if not os.path.exists(EXCEL_PATH):
    print(f"❌ Excel not found at {EXCEL_PATH}")
    sys.exit(1)

# === Categorize function ===
def categorize(details, model):
    d = (details or "").lower()
    m = (model or "").lower()
    if "fixed reader rfid" in d or "rfid" in d:
        return "RFID Reader", details
    if "mobile computer" in d or "mobile phone" in d or "rk26" in m:
        return "Mobile Computer", details
    if "tablet" in d:
        return "Tablet", details
    if "industrial barcode printer" in d or "industrial printer" in d:
        return "Industrial Printer", details
    if "mobile barcode printer" in d or "mobile printer" in d or "mobile receipt printer" in d:
        return "Mobile Printer", details
    if "receipt printer" in d:
        return "Receipt Printer", details
    if "desktop barcode printer" in d or "desktop printer" in d:
        return "Desktop Printer", details
    if "desktop barcode scanner" in d:
        return "Desktop Scanner", details
    if "industrial handheld" in d:
        return "Industrial Scanner", details
    if "handheld barcode scanner" in d or "handheld scanner" in d or "1d handheld" in d or "2d handheld" in d:
        return "Handheld Scanner", details
    if "barcode scanner" in d:
        return "Handheld Scanner", details
    return "Other", details

# === Parse status ===
def parse_status(remark, tag_id):
    r = (remark or "").strip()
    if not r or r == "*" or r == "":
        return "Available", None, False, None
    lower = r.lower()
    if "หาย" in lower or "lost" in lower or "สูญ" in lower:
        return "Lost", r, False, None
    if "ตัดให้" in lower or "ตัดขาย" in lower or "disposed" in lower:
        return "Disposed", r, False, None
    if (
        "ยืม" in lower
        or "ให้บ." in lower
        or "ให้พี่" in lower
        or "ใช้ระหว่าง" in lower
        or "ทดแทนงานซ่อม" in lower
    ):
        return "On_Loan", r, True, r
    return "Available", r, False, None

# === Default accessories ===
def default_accessories(category, brand, model):
    base = [f"เครื่อง {model}"]
    if category == "Mobile Computer":
        return json.dumps(base + ["แบตเตอรี่", "Power Adapter", "Power Cord", "USB Cable", "Cradle"], ensure_ascii=False)
    if category == "Tablet":
        return json.dumps(base + ["แบตเตอรี่", "Power Adapter", "Power Cord", "USB Cable"], ensure_ascii=False)
    if category in ("Handheld Scanner", "Desktop Scanner", "Industrial Scanner"):
        return json.dumps(base + ["USB Cable", "Stand"], ensure_ascii=False)
    if category in ("Desktop Printer", "Industrial Printer"):
        return json.dumps(base + ["Power Adapter", "Power Cord", "USB Cable", 'Ribbon Core 1"'], ensure_ascii=False)
    if category in ("Mobile Printer", "Receipt Printer"):
        return json.dumps(base + ["แบตเตอรี่", "Power Adapter", "Power Cord", "USB Cable"], ensure_ascii=False)
    if category == "RFID Reader":
        return json.dumps(base + ["แบตเตอรี่", "Power Adapter", "Power Cord", "USB Cable"], ensure_ascii=False)
    return json.dumps(base + ["Power Adapter", "Power Cord", "USB Cable"], ensure_ascii=False)

# === Read Excel ===
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb.active
print(f"📊 Sheet: {ws.title}, rows: {ws.max_row}, cols: {ws.max_column}")

parsed_rows = []
for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
    if row_idx == 1:
        continue  # header
    tag_id = str(row[0] or "").strip() if row[0] else ""
    details = str(row[1] or "").strip() if row[1] else ""
    model = str(row[2] or "").strip() if row[2] else ""
    serial_no = str(row[3] or "").strip() if row[3] else ""
    brand = str(row[4] or "").strip() if row[4] else ""
    remark = str(row[5] or "").strip() if row[5] else ""

    if not tag_id and not details and not model and not serial_no:
        continue
    if not tag_id and details:
        print(f"⚠️  Row {row_idx}: missing Tag ID, skipping (details: {details})")
        continue
    if not tag_id:
        continue

    # Clean serial — remove ".0" suffix
    if serial_no.endswith(".0"):
        serial_no = serial_no[:-2]

    parsed_rows.append({
        "tagId": tag_id,
        "details": details,
        "model": model,
        "serialNo": serial_no,
        "brand": brand,
        "remark": remark,
    })

print(f"📋 Parsed {len(parsed_rows)} valid rows")

# === Connect DB ===
conn = sqlite3.connect(DB_PATH)
conn.row_factory = sqlite3.Row
cur = conn.cursor()

# Find or create tenant
cur.execute("SELECT * FROM Tenant WHERE code = ?", (TENANT_CODE,))
tenant = cur.fetchone()
if not tenant:
    cur.execute(
        "INSERT INTO Tenant (id, name, code, plan, active, createdAt) VALUES (?, ?, ?, ?, 1, ?)",
        ("tenant-ais", "Active ID Solutions", TENANT_CODE, "enterprise", datetime.now().isoformat()),
    )
    conn.commit()
    cur.execute("SELECT * FROM Tenant WHERE code = ?", (TENANT_CODE,))
    tenant = cur.fetchone()
    print(f"✅ Created tenant: {tenant['code']}")
else:
    print(f"✅ Using existing tenant: {tenant['code']}")
tenant_id = tenant["id"]

# Find or create default location
cur.execute("SELECT * FROM Location WHERE tenantId = ? AND name = ?", (tenant_id, "HQ — คลังเดโม"))
location = cur.fetchone()
if not location:
    cur.execute(
        "INSERT INTO Location (id, tenantId, name, zone, createdAt) VALUES (?, ?, ?, ?, ?)",
        ("loc-hq", tenant_id, "HQ — คลังเดโม", "A", datetime.now().isoformat()),
    )
    conn.commit()
    cur.execute("SELECT * FROM Location WHERE tenantId = ? AND name = ?", (tenant_id, "HQ — คลังเดโม"))
    location = cur.fetchone()
    print(f"✅ Created location: {location['name']}")
else:
    print(f"✅ Using existing location: {location['name']}")
location_id = location["id"]

# Group by model+brand → create catalogs
catalog_map = {}
for row in parsed_rows:
    key = f"{row['model']}|{row['brand']}"
    if key not in catalog_map:
        cat, specs = categorize(row["details"], row["model"])
        catalog_map[key] = {"model": row["model"], "brand": row["brand"], "category": cat, "specs": specs}

print(f"📦 Unique catalogs: {len(catalog_map)}")

catalog_ids = {}
for key, cat in catalog_map.items():
    # Check if exists
    cur.execute(
        "SELECT * FROM ProductCatalog WHERE tenantId = ? AND model = ?",
        (tenant_id, cat["model"]),
    )
    existing = cur.fetchone()
    if existing:
        catalog_ids[key] = existing["id"]
        print(f"   ↳ Existing: {cat['brand']} {cat['model']} [{cat['category']}]")
    else:
        cat_id = f"cat-{len(catalog_ids)+1:04d}"
        cur.execute(
            """INSERT INTO ProductCatalog
               (id, tenantId, model, brand, category, defaultSpecs, defaultAccessories, createdAt, updatedAt)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                cat_id, tenant_id, cat["model"], cat["brand"], cat["category"],
                cat["specs"],
                default_accessories(cat["category"], cat["brand"], cat["model"]),
                datetime.now().isoformat(), datetime.now().isoformat(),
            ),
        )
        conn.commit()
        catalog_ids[key] = cat_id
        print(f"   ✨ Created: {cat['brand']} {cat['model']} [{cat['category']}]")

# Create asset units
created = 0
skipped = 0
on_loan = lost = disposed = available = 0

for row in parsed_rows:
    key = f"{row['model']}|{row['brand']}"
    catalog_id = catalog_ids.get(key)
    if not catalog_id:
        print(f"⚠️  No catalog for {row['tagId']}")
        skipped += 1
        continue

    # Check existing
    cur.execute(
        "SELECT * FROM AssetUnit WHERE tagId = ? OR (tenantId = ? AND serialNo = ?)",
        (row["tagId"], tenant_id, row["serialNo"]),
    )
    existing = cur.fetchone()
    if existing:
        print(f"   ↳ Existing asset: {row['tagId']}")
        skipped += 1
        continue

    status, notes, is_on_loan, loan_note = parse_status(row["remark"], row["tagId"])
    cat = catalog_map[key]
    accessories = default_accessories(cat["category"], row["brand"], row["model"])

    cur.execute(
        """INSERT INTO AssetUnit
           (tagId, tenantId, serialNo, catalogId, source, status, locationId,
            conditionGrade, registeredAccessories, notes, loanCount, createdAt, updatedAt)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (
            row["tagId"], tenant_id, row["serialNo"] or row["tagId"], catalog_id,
            "Demo_Stock", status, location_id,
            "A", accessories, notes or row["details"],
            0, datetime.now().isoformat(), datetime.now().isoformat(),
        ),
    )
    created += 1
    if status == "Available":
        available += 1
    elif status == "On_Loan":
        on_loan += 1
    elif status == "Lost":
        lost += 1
    elif status == "Disposed":
        disposed += 1

conn.commit()

# Print summary by category
print()
print("═══════════════════════════════════════════════")
print("🎉 Import Summary")
print("═══════════════════════════════════════════════")
print(f"Total parsed:    {len(parsed_rows)}")
print(f"Created assets:  {created}")
print(f"Skipped:         {skipped}")
print(f"  • Available:   {available}")
print(f"  • On_Loan:     {on_loan}  (Remark: พี่.../บ... ยืม/ให้...)")
print(f"  • Lost:        {lost}  (Remark: หาย)")
print(f"  • Disposed:    {disposed}  (Remark: ตัดให้...)")
print(f"Catalogs:        {len(catalog_ids)}")
print(f"Tenant:          {tenant['code']} ({tenant_id})")
print(f"Location:        {location['name']}")
print()

# Category breakdown
cur.execute(
    """SELECT pc.category, COUNT(*) as cnt
       FROM AssetUnit au
       JOIN ProductCatalog pc ON au.catalogId = pc.id
       WHERE au.tenantId = ?
       GROUP BY pc.category
       ORDER BY cnt DESC""",
    (tenant_id,),
)
print("📊 Assets by category:")
for r in cur.fetchall():
    print(f"   {r['category']:25s} {r['cnt']:4d}")

# Brand breakdown
cur.execute(
    """SELECT pc.brand, COUNT(*) as cnt
       FROM AssetUnit au
       JOIN ProductCatalog pc ON au.catalogId = pc.id
       WHERE au.tenantId = ?
       GROUP BY pc.brand
       ORDER BY cnt DESC""",
    (tenant_id,),
)
print()
print("📊 Assets by brand:")
for r in cur.fetchall():
    print(f"   {r['brand']:15s} {r['cnt']:4d}")

print()
print("═══════════════════════════════════════════════")

conn.close()
